import csv
import io
from pathlib import Path
from typing import List, Dict, Any
from pypdf import PdfReader
from docx import Document as DocxDocument

class DocumentParser:
    """Extract text from various document formats: PDF, DOCX, TXT, MD, CSV (REQ-IN-01)."""

    @staticmethod
    def parse(file_path: Path, file_type: str) -> str:
        ext = file_type.lower().lstrip(".")
        if ext == "pdf":
            return DocumentParser.parse_pdf(file_path)
        elif ext in ("docx", "doc"):
            return DocumentParser.parse_docx(file_path)
        elif ext in ("txt", "text", "md", "markdown"):
            return DocumentParser.parse_text(file_path)
        elif ext == "csv":
            return DocumentParser.parse_csv(file_path)
        else:
            # Fallback to UTF-8 text read
            try:
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    return f.read()
            except Exception as e:
                raise ValueError(f"Unsupported file format: {file_type}. Error: {str(e)}")

    @staticmethod
    def parse_pdf(file_path: Path) -> str:
        reader = PdfReader(str(file_path))
        extracted_pages = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                extracted_pages.append(f"[Page {i+1}]\n{text.strip()}")
        return "\n\n".join(extracted_pages)

    @staticmethod
    def parse_docx(file_path: Path) -> str:
        doc = DocxDocument(str(file_path))
        paragraphs = []
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    paragraphs.append(" | ".join(row_text))
        return "\n\n".join(paragraphs)

    @staticmethod
    def parse_text(file_path: Path) -> str:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    @staticmethod
    def parse_csv(file_path: Path) -> str:
        rows = []
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            header = None
            for i, row in enumerate(reader):
                if not row:
                    continue
                if i == 0:
                    header = row
                    rows.append("Headers: " + ", ".join(header))
                else:
                    if header and len(header) == len(row):
                        row_items = [f"{header[j]}: {val.strip()}" for j, val in enumerate(row) if val.strip()]
                        rows.append("Row: " + "; ".join(row_items))
                    else:
                        rows.append("Row: " + ", ".join(row))
        return "\n".join(rows)
